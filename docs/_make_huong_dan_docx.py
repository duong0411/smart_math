from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os


def set_run_font(run, size=11, bold=False, name='Times New Roman'):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.bold = bold


def add_heading_vn(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    size = 16 if level == 1 else (14 if level == 2 else 12)
    for run in p.runs:
        set_run_font(run, size=size, bold=True)
    return p


def add_para(doc, text, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(item)
        set_run_font(run)
        p.paragraph_format.space_after = Pt(2)


def add_code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=10, name='Consolas')
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(8)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, bold=True, size=10)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ''
            run = cell.paragraphs[0].add_run(val)
            set_run_font(run, size=10)
    doc.add_paragraph()


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run('HƯỚNG DẪN TẠO PHẦN MỀM EduSelf AI')
    set_run_font(r, size=18, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(
        '(Dùng cho hồ sơ dự thi sản phẩm sáng tạo — môn Toán / Địa lí)'
    )
    set_run_font(r, size=12)

    add_para(
        doc,
        'Tài liệu mô tả cách viết mã, kiểm thử, build và phát hành ứng dụng '
        'Flutter trên GitHub. Có thể in hoặc đưa vào thư mục hồ sơ dự thi.',
    )

    add_heading_vn(doc, '1. Sản phẩm là gì?', 1)
    add_para(
        doc,
        'EduSelf AI là ứng dụng học tập hỗ trợ học sinh THCS (lớp 6–9) '
        'bằng trí tuệ nhân tạo (Gemini):',
    )
    add_table(
        doc,
        ['Phiên bản', 'Nhánh GitHub', 'Nội dung chính'],
        [
            [
                'EduSelf Toán AI',
                'main',
                'Gia sư Toán, luyện tập, trò chơi toán, giám sát học tập',
            ],
            [
                'EduSelf Địa lí AI',
                'mon-dia-ly',
                'Gia sư Địa lí, luyện tập, khám phá địa lí, giám sát học tập',
            ],
        ],
    )
    add_para(doc, 'Repo: https://github.com/duong0411/smart_math')
    add_para(
        doc,
        'Ý tưởng dự thi: biến điện thoại / máy tính thành “gia sư AI” luôn '
        'sẵn sàng, bám chương trình GDPT Việt Nam, có game học tập và theo dõi '
        'tiến độ — không cần server riêng (API key lưu trên máy học sinh).',
    )

    add_heading_vn(doc, '2. Công nghệ đã dùng', 1)
    add_table(
        doc,
        ['Thành phần', 'Công nghệ', 'Vai trò'],
        [
            ['Ngôn ngữ', 'Dart', 'Viết logic ứng dụng'],
            ['Framework UI', 'Flutter', 'Một codebase → Android / iOS / Windows'],
            ['Quản lý trạng thái', 'Riverpod', 'Hồ sơ học sinh, API key, phiên chat…'],
            ['AI', 'Google Gemini API', 'Gia sư, tạo bài, chấm bài, báo cáo'],
            [
                'Lưu cục bộ',
                'SharedPreferences / secure storage',
                'Hồ sơ, lịch sử học, API key',
            ],
            [
                'Đọc tài liệu',
                'pdfrx (PDF), archive (Word .docx)',
                'Upload đề / bài làm',
            ],
            ['CI/CD', 'GitHub Actions', 'Tự build APK, IPA, EXE khi đẩy code'],
        ],
    )

    add_heading_vn(doc, '3. Các bước tạo phần mềm (từ ý tưởng → sản phẩm)', 1)
    add_code(
        doc,
        'Ý tưởng → Thiết kế chức năng → Cài môi trường → Viết code\n'
        '→ Kiểm thử (test) → Build thử trên máy → Đẩy lên GitHub\n'
        '→ GitHub Actions build APK / IPA / EXE → Tải file cài đặt',
    )

    add_heading_vn(doc, 'Bước 1 — Xác định chức năng', 2)
    add_bullets(
        doc,
        [
            'Gia sư AI — hỏi đáp từng bước, đính kèm ảnh / PDF / Word',
            'Luyện tập — AI tạo câu hỏi theo lớp, chấm bài',
            'Trò chơi — luyện kiến thức không cần API (offline)',
            'Giám sát — thống kê đúng/sai, gợi ý ôn',
            'Cài đặt — tên, lớp (6–9), Gemini API key',
        ],
    )

    add_heading_vn(doc, 'Bước 2 — Cài môi trường lập trình', 2)
    add_bullets(
        doc,
        [
            'Cài Flutter SDK (stable): https://docs.flutter.dev/get-started/install',
            'Cài Android Studio (build APK) và/hoặc Visual Studio (build Windows)',
            'Kiểm tra bằng lệnh: flutter doctor',
        ],
    )

    add_heading_vn(doc, 'Bước 3 — Lấy mã nguồn', 2)
    add_code(
        doc,
        'git clone https://github.com/duong0411/smart_math.git\n'
        'cd smart_math\n\n'
        '# Môn Toán\n'
        'git checkout main\n\n'
        '# hoặc môn Địa lí\n'
        'git checkout mon-dia-ly',
    )

    add_heading_vn(doc, 'Bước 4 — Cài thư viện & chạy thử', 2)
    add_code(doc, 'flutter pub get\nflutter run')
    add_para(
        doc,
        'Lấy API key tại Google AI Studio (https://aistudio.google.com/apikey) '
        '→ mở app → Cài đặt → dán key → Lưu.',
    )

    add_heading_vn(doc, '4. Cấu trúc mã nguồn', 1)
    add_code(
        doc,
        'smart_math/\n'
        '├── lib/\n'
        '│   ├── main.dart          # Điểm vào ứng dụng\n'
        '│   ├── core/              # Cấu hình, theme, router, AI\n'
        '│   ├── features/          # Gia sư, luyện tập, game…\n'
        '│   └── shared/            # Widget dùng chung\n'
        '├── assets/prompts/        # System prompt cho AI\n'
        '├── test/                  # Unit test\n'
        '├── .github/workflows/     # CI build APK + IPA + Windows\n'
        '└── pubspec.yaml',
    )
    add_para(
        doc,
        'lib/main.dart: khởi tạo Flutter → khởi tạo đọc PDF (pdfrx) → '
        'nạp cấu hình → chạy app với Riverpod.',
    )
    add_para(
        doc,
        'Luồng Gia sư: Màn hình UI → gọi Gemini AI → nhận câu trả lời → '
        'lưu local → hiển thị chat (tách lớp giao diện / logic / lưu trữ).',
    )

    add_heading_vn(doc, '5. Cách viết / chỉnh mã', 1)
    add_bullets(
        doc,
        [
            'Thêm màn hình: lib/features/.../presentation/pages/',
            'Đăng ký đường dẫn: lib/core/routing/app_router.dart',
            'Đổi tên app: lib/core/config/app_config.dart',
            'Đổi cách AI dạy: assets/prompts/eduself_system_prompt.md',
            'Thêm câu hỏi game: lib/features/math_games/',
        ],
    )
    add_code(doc, 'flutter analyze\nflutter test\nflutter run')

    add_heading_vn(doc, '6. Kiểm thử (Test code)', 1)
    add_code(doc, 'flutter test\nflutter analyze')
    add_para(
        doc,
        'Ý nghĩa trong hồ sơ dự thi: chứng minh sản phẩm có quy trình '
        'kiểm chứng chất lượng, không chỉ “chạy được”.',
    )

    add_heading_vn(doc, '7. Build trên máy cá nhân', 1)
    add_heading_vn(doc, 'Android APK', 2)
    add_code(doc, 'flutter build apk --release')
    add_para(doc, 'File ra: build/app/outputs/flutter-apk/app-release.apk')
    add_heading_vn(doc, 'Windows EXE', 2)
    add_code(
        doc,
        'flutter config --enable-windows-desktop\n'
        'flutter build windows --release',
    )
    add_para(
        doc,
        'Chạy eduself_study_app.exe trong build/windows/x64/runner/Release/ '
        '(giữ nguyên cả thư mục DLL + data).',
    )
    add_heading_vn(doc, 'iOS IPA', 2)
    add_para(
        doc,
        'Cần máy Mac + Xcode. Trên GitHub Actions sẽ đóng gói IPA '
        '(unsigned — chỉ sideload / kiểm thử).',
    )

    add_heading_vn(doc, '8. Build tự động trên GitHub', 1)
    add_para(doc, 'File cấu hình: .github/workflows/build-mobile.yml')
    add_para(
        doc,
        'Khi push code lên nhánh đúng (hoặc bấm Run workflow), GitHub chạy 3 job:',
    )
    add_table(
        doc,
        ['Job', 'Máy ảo CI', 'Kết quả'],
        [
            ['Android APK', 'ubuntu-latest', 'app-release.apk'],
            ['iOS IPA', 'macos-latest', 'SmartMath.ipa'],
            ['Windows EXE', 'windows-latest', 'SmartMath-windows.zip'],
        ],
    )
    add_heading_vn(doc, 'Các bước', 2)
    add_bullets(
        doc,
        [
            'Sửa code → git commit → git push',
            'Vào repo GitHub → tab Actions',
            'Xem workflow Build APK, iOS & Windows (các job xanh)',
            'Tải Artifacts hoặc Releases',
        ],
    )
    add_table(
        doc,
        ['Sản phẩm', 'Nhánh', 'Tag release (ví dụ)'],
        [
            ['Toán', 'main', 'android-v1.0.* / ios-v1.0.* / windows-v1.0.*'],
            [
                'Địa lí',
                'mon-dia-ly',
                'dialy-android-v1.0.* / dialy-ios-v1.0.* / dialy-windows-v1.0.*',
            ],
        ],
    )

    add_heading_vn(doc, '9. Quy trình Git gợi ý', 1)
    add_code(
        doc,
        'git status\n'
        'git add .\n'
        'git commit -m "mô tả ngắn thay đổi"\n'
        'git push origin main          # Toán\n'
        '# hoặc\n'
        'git push origin mon-dia-ly    # Địa lí',
    )
    add_para(doc, 'Lưu ý: không commit API key, file .env thật, hay thư mục build/.')

    add_heading_vn(doc, '10. Checklist hồ sơ dự thi', 1)
    add_bullets(
        doc,
        [
            'Mô tả ý tưởng & đối tượng (học sinh THCS lớp 6–9)',
            'Ảnh / video demo các chức năng chính',
            'Link GitHub mã nguồn',
            'Ảnh màn hình Actions (3 job build thành công)',
            'Link Releases tải APK / Windows / IPA',
            'Bảng công nghệ (mục 2)',
            'Ảnh kết quả lệnh flutter test',
            'Hướng dẫn giáo viên / giám khảo cài app thử',
        ],
    )

    add_heading_vn(doc, '11. Tóm tắt một dòng', 1)
    add_para(
        doc,
        'Viết app bằng Flutter → kiểm thử bằng flutter test → đẩy GitHub → '
        'GitHub Actions tự build APK, IPA và EXE → tải về dùng và đưa vào hồ sơ dự thi.',
        bold=True,
    )

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run(
        '\n— Tài liệu kỹ thuật đi kèm mã nguồn EduSelf / Smart Math —'
    )
    set_run_font(r, size=10)
    return doc


def main():
    doc = build_doc()
    paths = [
        r'C:\Users\Duong Phung\Downloads\eduself\docs\HUONG_DAN_TAO_PHAN_MEM.docx',
        r'C:\Users\Duong Phung\Downloads\smart_math\smart_math\docs\HUONG_DAN_TAO_PHAN_MEM.docx',
        r'C:\Users\Duong Phung\Downloads\HUONG_DAN_TAO_PHAN_MEM.docx',
    ]
    for path in paths:
        os.makedirs(os.path.dirname(path) or '.', exist_ok=True)
        doc.save(path)
        print('Saved', path)


if __name__ == '__main__':
    main()
